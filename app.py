from flask import Flask, render_template, request, send_file
from docx import Document
import os
import tempfile

app = Flask(__name__)

# Define the path to the DOCX file
DOCX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Advanced_java_practical_file[1].docx')

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/process', methods=['POST'])
def process():
    name = request.form['name']
    roll_no = request.form['roll_no']
    
    try:
        # Open the original document
        doc = Document(DOCX_PATH)
        
        # Names to search for (including variations)
        name_variations = ['Aryaman Sharma', 'Aryaman sharma', 'ARYAMAN SHARMA', 'Aryaman sharman']
        roll_variations = ['22011403007', '22011403027']
        
        # Process each paragraph
        for paragraph in doc.paragraphs:
            # Store the original text
            original_text = paragraph.text
            
            # Check for name variations
            for old_name in name_variations:
                if old_name in original_text:
                    # Replace text in each run to preserve formatting
                    for run in paragraph.runs:
                        run.text = run.text.replace(old_name, name)
            
            # Check for roll number variations
            for old_roll in roll_variations:
                if old_roll in original_text:
                    # Replace text in each run to preserve formatting
                    for run in paragraph.runs:
                        run.text = run.text.replace(old_roll, roll_no)
        
        # Also check text in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        # Check for name variations
                        for old_name in name_variations:
                            if old_name in paragraph.text:
                                for run in paragraph.runs:
                                    run.text = run.text.replace(old_name, name)
                        
                        # Check for roll number variations
                        for old_roll in roll_variations:
                            if old_roll in paragraph.text:
                                for run in paragraph.runs:
                                    run.text = run.text.replace(old_roll, roll_no)
        
        # Create a temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp:
            temp_path = tmp.name
            doc.save(temp_path)
            
            # Send the file and then delete it
            response = send_file(
                temp_path,
                as_attachment=True,
                download_name='modified_document.docx',
                mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            
            # Delete the temp file after sending
            @response.call_on_close
            def cleanup():
                if os.path.exists(temp_path):
                    os.remove(temp_path)
                    
            return response

    except Exception as e:
        return f"Error processing the document: {str(e)}", 500

if __name__ == '__main__':
    app.run(debug=True) 